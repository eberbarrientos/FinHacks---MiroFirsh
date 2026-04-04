#!/usr/bin/env python3
"""
Generate a stylish PowerPoint presentation explaining the
Climate Risk Intelligence Platform in layman terms.

Usage:
    pip install python-pptx
    python create_presentation.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── Colour palette ──────────────────────────────────────────────────────
DARK_BG       = RGBColor(0x0F, 0x17, 0x2A)   # Deep navy
CARD_BG       = RGBColor(0x1E, 0x29, 0x3B)   # Slate card
ACCENT_CYAN   = RGBColor(0x06, 0xB6, 0xD4)   # Cyan-500
ACCENT_TEAL   = RGBColor(0x14, 0xB8, 0xA6)   # Teal-500
ACCENT_AMBER  = RGBColor(0xF5, 0x9E, 0x0B)   # Amber-500
ACCENT_RED    = RGBColor(0xEF, 0x44, 0x44)   # Red-500
ACCENT_GREEN  = RGBColor(0x10, 0xB9, 0x81)   # Emerald-500
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY    = RGBColor(0xCB, 0xD5, 0xE1)   # Slate-300
MID_GRAY      = RGBColor(0x94, 0xA3, 0xB8)   # Slate-400
DIM_GRAY      = RGBColor(0x64, 0x74, 0x8B)   # Slate-500
GRADIENT_MID  = RGBColor(0x0E, 0x24, 0x3D)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ── Helpers ─────────────────────────────────────────────────────────────

def set_slide_bg(slide, color):
    """Set solid background colour for a slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color, corner_radius=Inches(0.15)):
    """Add a rounded rectangle shape."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    shape.shadow.inherit = False
    # Adjust corner rounding
    if hasattr(shape, 'adjustments') and len(shape.adjustments) > 0:
        shape.adjustments[0] = 0.05
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=14,
                 color=WHITE, bold=False, alignment=PP_ALIGN.LEFT,
                 font_name="Calibri"):
    """Add a text box with styled text."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_multiline_text(slide, left, top, width, height, lines,
                       font_size=14, color=WHITE, line_spacing=1.3,
                       font_name="Calibri"):
    """Add a text box with multiple styled lines (list of dicts)."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line_info in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line_info.get("text", "")
        p.font.size = Pt(line_info.get("size", font_size))
        p.font.color.rgb = line_info.get("color", color)
        p.font.bold = line_info.get("bold", False)
        p.font.name = line_info.get("font", font_name)
        p.alignment = line_info.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(line_info.get("space_after", 4))
        if "space_before" in line_info:
            p.space_before = Pt(line_info["space_before"])
    return txBox


def add_accent_line(slide, left, top, width, color=ACCENT_CYAN, height=Inches(0.04)):
    """Add a thin accent line."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_icon_circle(slide, left, top, size, fill_color, text, text_size=18):
    """Add a circle with an emoji/icon character."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left, top, size, size
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(text_size)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return shape


def add_card(slide, left, top, width, height, title, body_lines,
             accent_color=ACCENT_CYAN, icon_text="", title_size=16, body_size=12):
    """Add a styled card with title, accent bar, and body text."""
    card = add_shape(slide, left, top, width, height, CARD_BG)
    # Accent bar at top
    add_accent_line(slide, left + Inches(0.2), top + Inches(0.12),
                    Inches(0.6), accent_color, Inches(0.035))
    # Icon
    if icon_text:
        add_text_box(slide, left + Inches(0.2), top + Inches(0.22),
                     Inches(0.5), Inches(0.4), icon_text,
                     font_size=20, color=accent_color, bold=True)
    # Title
    title_left = left + (Inches(0.7) if icon_text else Inches(0.2))
    add_text_box(slide, title_left, top + Inches(0.22),
                 width - Inches(0.5), Inches(0.4), title,
                 font_size=title_size, color=WHITE, bold=True)
    # Body
    y = top + Inches(0.7)
    for line in body_lines:
        add_text_box(slide, left + Inches(0.25), y,
                     width - Inches(0.5), Inches(0.35), line,
                     font_size=body_size, color=LIGHT_GRAY)
        y += Inches(0.32)
    return card


def add_numbered_step(slide, left, top, number, title, desc,
                      accent_color=ACCENT_CYAN, width=Inches(11)):
    """Add a numbered step row."""
    # Number circle
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left, top, Inches(0.55), Inches(0.55)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = accent_color
    circle.line.fill.background()
    tf = circle.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = str(number)
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Title
    add_text_box(slide, left + Inches(0.75), top - Inches(0.02),
                 Inches(3), Inches(0.35), title,
                 font_size=16, color=WHITE, bold=True)
    # Description
    add_text_box(slide, left + Inches(0.75), top + Inches(0.28),
                 width - Inches(1), Inches(0.35), desc,
                 font_size=12, color=LIGHT_GRAY)


# ── Slide builders ──────────────────────────────────────────────────────

def build_title_slide(prs):
    """Slide 1: Title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    set_slide_bg(slide, DARK_BG)

    # Decorative top bar
    add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_CYAN, Inches(0.06))

    # Decorative side accent
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.06), Inches(0.08), Inches(7.44)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_TEAL
    shape.line.fill.background()

    # Globe icon area
    globe = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(5.5), Inches(1.5), Inches(2.2), Inches(2.2)
    )
    globe.fill.solid()
    globe.fill.fore_color.rgb = RGBColor(0x06, 0x4E, 0x5B)
    globe.line.fill.background()
    tf = globe.text_frame
    p = tf.paragraphs[0]
    p.text = "🌍"
    p.font.size = Pt(60)
    p.alignment = PP_ALIGN.CENTER

    # Title
    add_multiline_text(slide, Inches(0.8), Inches(1.6), Inches(5), Inches(2.5), [
        {"text": "Climate Risk", "size": 44, "color": WHITE, "bold": True, "space_after": 2},
        {"text": "Intelligence Platform", "size": 44, "color": ACCENT_CYAN, "bold": True, "space_after": 12},
        {"text": "Agent-Based Cascade Simulation for Portfolio Risk", "size": 18,
         "color": LIGHT_GRAY, "space_after": 6},
        {"text": "Powered by MiroFish Architecture  •  Graph RAG  •  LLM Agents",
         "size": 13, "color": MID_GRAY, "space_after": 0},
    ])

    # Accent line under title
    add_accent_line(slide, Inches(0.8), Inches(4.3), Inches(4), ACCENT_CYAN)

    # Bottom tagline
    add_text_box(slide, Inches(0.8), Inches(5.2), Inches(8), Inches(0.8),
                 "Describe any climate event in plain English — the platform simulates\n"
                 "how damage cascades through company dependencies and your portfolio.",
                 font_size=14, color=MID_GRAY)

    # Bottom bar
    add_accent_line(slide, Inches(0), Inches(7.44), SLIDE_W, ACCENT_TEAL, Inches(0.06))


def build_problem_slide(prs):
    """Slide 2: The Problem — Why climate risk matters."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_RED, Inches(0.04))

    # Section label
    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(3), Inches(0.3),
                 "THE CHALLENGE", font_size=11, color=ACCENT_RED, bold=True)

    add_text_box(slide, Inches(0.8), Inches(0.7), Inches(10), Inches(0.6),
                 "Climate Events Create Hidden Financial Risks",
                 font_size=32, color=WHITE, bold=True)

    add_text_box(slide, Inches(0.8), Inches(1.35), Inches(10), Inches(0.5),
                 "A single hurricane doesn't just damage one company — it triggers a chain reaction "
                 "through supply chains, insurance networks, and infrastructure dependencies.",
                 font_size=14, color=LIGHT_GRAY)

    # Three problem cards
    cards = [
        {
            "icon": "🌀", "title": "Direct Damage",
            "color": ACCENT_RED,
            "lines": [
                "A hurricane hits Texas and",
                "damages oil refineries, power",
                "grids, and coastal real estate.",
                "This is the obvious loss.",
            ]
        },
        {
            "icon": "🔗", "title": "Hidden Cascades",
            "color": ACCENT_AMBER,
            "lines": [
                "When the power grid fails,",
                "factories shut down. When",
                "refineries close, fuel prices",
                "spike across the country.",
            ]
        },
        {
            "icon": "📉", "title": "Portfolio Impact",
            "color": RGBColor(0xA7, 0x8B, 0xFA),
            "lines": [
                "Your investments in seemingly",
                "unrelated companies suffer",
                "because they depend on the",
                "same damaged infrastructure.",
            ]
        },
    ]

    x_start = Inches(0.8)
    card_w = Inches(3.7)
    gap = Inches(0.35)
    for i, c in enumerate(cards):
        x = x_start + i * (card_w + gap)
        add_card(slide, x, Inches(2.2), card_w, Inches(2.8),
                 c["title"], c["lines"], c["color"], c["icon"])

    # Bottom insight box
    insight_shape = add_shape(slide, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.5),
                              RGBColor(0x1A, 0x1A, 0x2E))
    add_text_box(slide, Inches(1.1), Inches(5.55), Inches(1), Inches(0.3),
                 "💡", font_size=20, color=ACCENT_AMBER)
    add_multiline_text(slide, Inches(1.6), Inches(5.55), Inches(10.5), Inches(1.2), [
        {"text": "The Key Insight", "size": 16, "color": ACCENT_AMBER, "bold": True, "space_after": 6},
        {"text": "Traditional risk models look at each company in isolation. But in reality, companies are "
                 "deeply interconnected. A single climate event can amplify losses 2-5x through these hidden "
                 "dependency chains. Our platform maps and simulates these cascades.",
         "size": 13, "color": LIGHT_GRAY, "space_after": 0},
    ])


def build_solution_overview_slide(prs):
    """Slide 3: What the platform does — high-level."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_CYAN, Inches(0.04))

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(3), Inches(0.3),
                 "THE SOLUTION", font_size=11, color=ACCENT_CYAN, bold=True)

    add_text_box(slide, Inches(0.8), Inches(0.7), Inches(10), Inches(0.6),
                 "Describe an Event. See the Cascade. Protect Your Portfolio.",
                 font_size=30, color=WHITE, bold=True)

    add_text_box(slide, Inches(0.8), Inches(1.3), Inches(10), Inches(0.5),
                 "Type any climate scenario in plain English and watch AI agents simulate "
                 "how damage ripples through the economy.",
                 font_size=14, color=LIGHT_GRAY)

    # Example input box
    input_shape = add_shape(slide, Inches(0.8), Inches(2.1), Inches(11.7), Inches(0.7),
                            RGBColor(0x0F, 0x23, 0x3D))
    add_text_box(slide, Inches(1.0), Inches(2.2), Inches(0.8), Inches(0.4),
                 "🔍", font_size=18, color=MID_GRAY)
    add_text_box(slide, Inches(1.6), Inches(2.22), Inches(10), Inches(0.4),
                 '"Category 5 hurricane hits the Texas Gulf Coast with 160mph winds"',
                 font_size=15, color=ACCENT_CYAN, bold=False)

    # Quick scenario pills
    scenarios = ["🌀 Hurricane", "🔥 Wildfire", "🌊 Flood", "🌡️ Heatwave",
                 "💰 Carbon Tax", "🏜️ Drought", "🌪️ Tornado", "❄️ Winter Storm"]
    x = Inches(0.8)
    for s in scenarios:
        pill = add_shape(slide, x, Inches(3.0), Inches(1.35), Inches(0.38),
                         RGBColor(0x16, 0x2D, 0x4A))
        add_text_box(slide, x + Inches(0.08), Inches(3.02),
                     Inches(1.3), Inches(0.35), s,
                     font_size=10, color=LIGHT_GRAY, alignment=PP_ALIGN.CENTER)
        x += Inches(1.45)

    # What you get — 5 output cards
    outputs = [
        {"icon": "🕸️", "title": "Agent Network", "desc": "Interactive graph of companies\nand their dependencies",
         "color": ACCENT_CYAN},
        {"icon": "⏱️", "title": "Cascade Timeline", "desc": "Round-by-round view of\nhow damage spreads",
         "color": ACCENT_TEAL},
        {"icon": "📝", "title": "AI Narrative", "desc": "Executive briefing explaining\nwhat happened and why",
         "color": ACCENT_GREEN},
        {"icon": "📊", "title": "Loss Analysis", "desc": "Sortable table of every\naffected company and loss",
         "color": ACCENT_AMBER},
        {"icon": "🛡️", "title": "Recommendations", "desc": "Actionable steps to protect\nyour portfolio",
         "color": RGBColor(0xA7, 0x8B, 0xFA)},
    ]

    card_w = Inches(2.15)
    gap = Inches(0.2)
    x_start = Inches(0.8)
    for i, o in enumerate(outputs):
        x = x_start + i * (card_w + gap)
        card = add_shape(slide, x, Inches(3.8), card_w, Inches(2.6), CARD_BG)
        add_accent_line(slide, x + Inches(0.15), Inches(3.92),
                        Inches(0.5), o["color"], Inches(0.03))
        add_text_box(slide, x + Inches(0.15), Inches(4.05),
                     Inches(0.5), Inches(0.4), o["icon"],
                     font_size=22, color=o["color"])
        add_text_box(slide, x + Inches(0.15), Inches(4.55),
                     card_w - Inches(0.3), Inches(0.35), o["title"],
                     font_size=14, color=WHITE, bold=True)
        add_text_box(slide, x + Inches(0.15), Inches(4.95),
                     card_w - Inches(0.3), Inches(0.8), o["desc"],
                     font_size=11, color=MID_GRAY)

    # Bottom stat bar
    stats_shape = add_shape(slide, Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.55),
                            RGBColor(0x0F, 0x23, 0x3D))
    stat_items = [
        ("200+", "Real Companies"),
        ("12", "Industry Sectors"),
        ("50+", "US Regions"),
        ("3-5", "Cascade Rounds"),
        ("15-40", "AI Agents per Sim"),
    ]
    sx = Inches(1.2)
    for val, label in stat_items:
        add_text_box(slide, sx, Inches(6.73), Inches(1), Inches(0.25),
                     val, font_size=16, color=ACCENT_CYAN, bold=True,
                     alignment=PP_ALIGN.CENTER)
        add_text_box(slide, sx, Inches(6.98), Inches(1), Inches(0.25),
                     label, font_size=9, color=MID_GRAY,
                     alignment=PP_ALIGN.CENTER)
        sx += Inches(2.2)


def build_how_it_works_slide(prs):
    """Slide 4: How it works — the 6-step pipeline."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_TEAL, Inches(0.04))

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(3), Inches(0.3),
                 "HOW IT WORKS", font_size=11, color=ACCENT_TEAL, bold=True)

    add_text_box(slide, Inches(0.8), Inches(0.7), Inches(10), Inches(0.6),
                 "From Plain English to Portfolio Insights in 6 Steps",
                 font_size=30, color=WHITE, bold=True)

    steps = [
        ("Event Parsing",
         "You type a climate event in plain English. AI understands the disaster type, affected regions, and impacted industries.",
         ACCENT_CYAN),
        ("Agent Generation",
         "The system creates AI agents for real companies relevant to the disaster — ExxonMobil for Texas hurricanes, PG&E for California wildfires.",
         ACCENT_TEAL),
        ("Knowledge Graph",
         "A dependency network is built: who supplies whom, who insures whom, who powers whom. This maps the hidden connections.",
         ACCENT_GREEN),
        ("Cascade Simulation",
         "Round 0: direct damage hits. Rounds 1-3: each AI agent reasons about its own losses based on damaged dependencies (Graph RAG).",
         ACCENT_AMBER),
        ("Narrative & Analysis",
         "AI writes an executive briefing explaining the cascade, quantifies losses per company, and shows amplification effects.",
         RGBColor(0xA7, 0x8B, 0xFA)),
        ("Recommendations",
         "Actionable portfolio advice: rebalance, hedge, diversify, insure, or exit positions based on simulation results.",
         ACCENT_RED),
    ]

    y = Inches(1.6)
    for i, (title, desc, color) in enumerate(steps):
        add_numbered_step(slide, Inches(0.8), y, i + 1, title, desc, color)
        y += Inches(0.88)

    # Connecting line on the left
    line_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.05), Inches(2.1), Inches(0.02), Inches(4.5)
    )
    line_shape.fill.solid()
    line_shape.fill.fore_color.rgb = DIM_GRAY
    line_shape.line.fill.background()


def build_architecture_slide(prs):
    """Slide 5: Architecture — tech stack and data flow."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, RGBColor(0xA7, 0x8B, 0xFA), Inches(0.04))

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(3), Inches(0.3),
                 "ARCHITECTURE", font_size=11, color=RGBColor(0xA7, 0x8B, 0xFA), bold=True)

    add_text_box(slide, Inches(0.8), Inches(0.7), Inches(10), Inches(0.6),
                 "Full-Stack Platform with AI-Powered Simulation Engine",
                 font_size=30, color=WHITE, bold=True)

    # Left column: Frontend
    add_card(slide, Inches(0.5), Inches(1.6), Inches(3.8), Inches(3.2),
             "🖥️  Frontend", [
                 "Next.js 15 + TypeScript",
                 "Tailwind CSS + shadcn/ui",
                 "Canvas-based network graph",
                 "Framer Motion animations",
                 "Zustand state management",
                 "Real-time simulation display",
             ], ACCENT_CYAN, title_size=18, body_size=12)

    # Middle column: Backend
    add_card(slide, Inches(4.6), Inches(1.6), Inches(3.8), Inches(3.2),
             "⚙️  Backend", [
                 "FastAPI (Python 3.11+)",
                 "SQLAlchemy ORM + Alembic",
                 "PostgreSQL / SQLite",
                 "Redis caching layer",
                 "Pydantic validation",
                 "RESTful API endpoints",
             ], ACCENT_TEAL, title_size=18, body_size=12)

    # Right column: AI Engine
    add_card(slide, Inches(8.7), Inches(1.6), Inches(4.1), Inches(3.2),
             "🧠  AI Engine", [
                 "Google Gemini / OpenAI LLM",
                 "NetworkX knowledge graph",
                 "Graph RAG context retrieval",
                 "Per-agent LLM reasoning",
                 "Rule-based fallback mode",
                 "Rate-limited API management",
             ], RGBColor(0xA7, 0x8B, 0xFA), title_size=18, body_size=12)

    # Data flow arrow section
    flow_y = Inches(5.1)
    add_text_box(slide, Inches(0.8), flow_y, Inches(3), Inches(0.3),
                 "DATA FLOW", font_size=11, color=ACCENT_AMBER, bold=True)

    flow_steps = [
        ("User Input", ACCENT_CYAN),
        ("Event Parser", ACCENT_TEAL),
        ("Agent Builder", ACCENT_GREEN),
        ("Knowledge Graph", RGBColor(0xA7, 0x8B, 0xFA)),
        ("Cascade Engine", ACCENT_AMBER),
        ("Results + UI", ACCENT_RED),
    ]

    x = Inches(0.5)
    box_w = Inches(1.7)
    arrow_w = Inches(0.35)
    for i, (label, color) in enumerate(flow_steps):
        box = add_shape(slide, x, flow_y + Inches(0.4), box_w, Inches(0.5), CARD_BG)
        add_accent_line(slide, x, flow_y + Inches(0.4), box_w, color, Inches(0.03))
        add_text_box(slide, x + Inches(0.05), flow_y + Inches(0.45),
                     box_w - Inches(0.1), Inches(0.4), label,
                     font_size=11, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
        if i < len(flow_steps) - 1:
            add_text_box(slide, x + box_w, flow_y + Inches(0.42),
                         arrow_w, Inches(0.4), "→",
                         font_size=18, color=DIM_GRAY, alignment=PP_ALIGN.CENTER)
        x += box_w + arrow_w

    # API endpoints
    api_y = Inches(6.3)
    add_text_box(slide, Inches(0.8), api_y, Inches(3), Inches(0.3),
                 "KEY API ENDPOINTS", font_size=11, color=MID_GRAY, bold=True)
    endpoints = [
        "POST /api/simulate/run — Run cascade simulation",
        "GET  /api/portfolios — List portfolios",
        "GET  /api/portfolios/{id} — Portfolio details",
        "GET  /health — Service health check",
    ]
    for i, ep in enumerate(endpoints):
        add_text_box(slide, Inches(0.8) + (i % 2) * Inches(6), api_y + Inches(0.3) + (i // 2) * Inches(0.28),
                     Inches(5.5), Inches(0.25), ep,
                     font_size=10, color=DIM_GRAY, font_name="Courier New")


def build_agents_slide(prs):
    """Slide 6: AI Agents & Graph RAG — the MiroFish magic."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_GREEN, Inches(0.04))

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(3), Inches(0.3),
                 "AI AGENTS & GRAPH RAG", font_size=11, color=ACCENT_GREEN, bold=True)

    add_text_box(slide, Inches(0.8), Inches(0.7), Inches(10), Inches(0.6),
                 "Every Company Thinks for Itself",
                 font_size=30, color=WHITE, bold=True)

    add_text_box(slide, Inches(0.8), Inches(1.3), Inches(10), Inches(0.5),
                 "Inspired by MiroFish — each company becomes an autonomous AI agent that reasons "
                 "about its own situation using its neighborhood in the dependency graph.",
                 font_size=14, color=LIGHT_GRAY)

    # Left: Agent anatomy
    add_text_box(slide, Inches(0.8), Inches(2.1), Inches(3), Inches(0.3),
                 "AGENT ANATOMY", font_size=11, color=ACCENT_CYAN, bold=True)

    agent_props = [
        ("🏢", "Identity", "Company name, sector, region, type"),
        ("💰", "Financials", "Market value, carbon intensity"),
        ("🔗", "Dependencies", "Supply chain, grid, insurance links"),
        ("🧠", "Reasoning", "LLM-powered or rule-based logic"),
        ("📊", "Graph Context", "2-hop neighborhood + damage state"),
        ("💬", "Social Feed", "Other agents' reactions this round"),
    ]

    y = Inches(2.5)
    for icon, title, desc in agent_props:
        row_shape = add_shape(slide, Inches(0.8), y, Inches(5.2), Inches(0.48), CARD_BG)
        add_text_box(slide, Inches(0.95), y + Inches(0.04),
                     Inches(0.4), Inches(0.35), icon, font_size=14, color=ACCENT_CYAN)
        add_text_box(slide, Inches(1.4), y + Inches(0.04),
                     Inches(1.5), Inches(0.35), title,
                     font_size=12, color=WHITE, bold=True)
        add_text_box(slide, Inches(2.8), y + Inches(0.04),
                     Inches(3), Inches(0.35), desc,
                     font_size=11, color=MID_GRAY)
        y += Inches(0.55)

    # Right: Graph RAG explanation
    add_text_box(slide, Inches(6.5), Inches(2.1), Inches(3), Inches(0.3),
                 "GRAPH RAG PATTERN", font_size=11, color=RGBColor(0xA7, 0x8B, 0xFA), bold=True)

    graph_card = add_shape(slide, Inches(6.5), Inches(2.5), Inches(6), Inches(4.2), CARD_BG)

    graph_steps = [
        {"text": "1. Build Knowledge Graph", "size": 14, "color": ACCENT_CYAN, "bold": True, "space_after": 2},
        {"text": "   Companies become nodes. Dependencies become edges.", "size": 12, "color": LIGHT_GRAY, "space_after": 10},
        {"text": "2. Walk the Neighborhood", "size": 14, "color": ACCENT_TEAL, "bold": True, "space_after": 2},
        {"text": "   For each agent, retrieve 2-hop neighbors + damage state.", "size": 12, "color": LIGHT_GRAY, "space_after": 10},
        {"text": "3. Augment the Prompt", "size": 14, "color": ACCENT_GREEN, "bold": True, "space_after": 2},
        {"text": "   Inject graph context into the LLM system prompt.", "size": 12, "color": LIGHT_GRAY, "space_after": 10},
        {"text": "4. Agent Reasons", "size": 14, "color": ACCENT_AMBER, "bold": True, "space_after": 2},
        {"text": "   LLM sees damaged dependencies and estimates its own loss.", "size": 12, "color": LIGHT_GRAY, "space_after": 10},
        {"text": "5. Interactive Cascade", "size": 14, "color": RGBColor(0xA7, 0x8B, 0xFA), "bold": True, "space_after": 2},
        {"text": "   Agent A's reaction feeds into Agent B's reasoning.", "size": 12, "color": LIGHT_GRAY, "space_after": 10},
        {"text": "6. Emergent Behavior", "size": 14, "color": ACCENT_RED, "bold": True, "space_after": 2},
        {"text": "   Complex cascade patterns emerge from simple agent rules.", "size": 12, "color": LIGHT_GRAY, "space_after": 0},
    ]

    add_multiline_text(slide, Inches(6.8), Inches(2.7), Inches(5.4), Inches(3.8), graph_steps)

    # Bottom: MiroFish comparison
    mf_shape = add_shape(slide, Inches(0.5), Inches(6.2), Inches(12.3), Inches(1.0),
                         RGBColor(0x0F, 0x23, 0x3D))
    add_text_box(slide, Inches(0.8), Inches(6.3), Inches(2), Inches(0.3),
                 "🐟 MIROFISH PATTERN", font_size=11, color=ACCENT_TEAL, bold=True)

    comparisons = [
        ("Zep Knowledge Graph", "NetworkX DiGraph"),
        ("Zep graph search", "2-hop neighborhood walk"),
        ("Per-agent LLM persona", "Company persona + graph RAG"),
        ("OASIS social rounds", "Cascade rounds with propagation"),
    ]
    cx = Inches(0.8)
    for orig, ours in comparisons:
        add_text_box(slide, cx, Inches(6.6), Inches(2.8), Inches(0.22),
                     f"{orig}  →  {ours}",
                     font_size=10, color=MID_GRAY)
        cx += Inches(3.0)


def build_cascade_example_slide(prs):
    """Slide 7: Cascade example — Texas hurricane walkthrough."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_AMBER, Inches(0.04))

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(3), Inches(0.3),
                 "EXAMPLE WALKTHROUGH", font_size=11, color=ACCENT_AMBER, bold=True)

    add_text_box(slide, Inches(0.8), Inches(0.7), Inches(10), Inches(0.6),
                 '"Category 5 Hurricane Hits the Texas Gulf Coast"',
                 font_size=28, color=WHITE, bold=True)

    # Round 0
    r0_shape = add_shape(slide, Inches(0.5), Inches(1.6), Inches(3.8), Inches(2.8), CARD_BG)
    add_accent_line(slide, Inches(0.5), Inches(1.6), Inches(3.8), ACCENT_RED, Inches(0.04))
    add_multiline_text(slide, Inches(0.7), Inches(1.75), Inches(3.4), Inches(2.5), [
        {"text": "Round 0 — Direct Impact", "size": 16, "color": ACCENT_RED, "bold": True, "space_after": 8},
        {"text": "🏭 ExxonMobil — refinery damage", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "⚡ CenterPoint Energy — grid failure", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "🏗️ Valero Energy — operations halted", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "🏠 Camden Property — property damage", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "✈️ Southwest Airlines — flights cancelled", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "⚡ ERCOT — grid operator overwhelmed", "size": 12, "color": LIGHT_GRAY, "space_after": 0},
    ])

    # Arrow
    add_text_box(slide, Inches(4.35), Inches(2.7), Inches(0.5), Inches(0.4),
                 "→", font_size=28, color=ACCENT_AMBER)

    # Round 1
    r1_shape = add_shape(slide, Inches(4.8), Inches(1.6), Inches(3.8), Inches(2.8), CARD_BG)
    add_accent_line(slide, Inches(4.8), Inches(1.6), Inches(3.8), ACCENT_AMBER, Inches(0.04))
    add_multiline_text(slide, Inches(5.0), Inches(1.75), Inches(3.4), Inches(2.5), [
        {"text": "Round 1 — First Cascade", "size": 16, "color": ACCENT_AMBER, "bold": True, "space_after": 8},
        {"text": "🔗 Phillips 66 — supply chain disrupted", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "🏦 Chubb Insurance — claims surge", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "🚂 BNSF Railway — routes blocked", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "🧪 LyondellBasell — no power/feedstock", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "💻 Dell Technologies — factory offline", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "Each agent sees damaged deps & reacts", "size": 11, "color": DIM_GRAY, "space_after": 0},
    ])

    # Arrow
    add_text_box(slide, Inches(8.65), Inches(2.7), Inches(0.5), Inches(0.4),
                 "→", font_size=28, color=ACCENT_AMBER)

    # Round 2+
    r2_shape = add_shape(slide, Inches(9.1), Inches(1.6), Inches(3.8), Inches(2.8), CARD_BG)
    add_accent_line(slide, Inches(9.1), Inches(1.6), Inches(3.8), RGBColor(0xA7, 0x8B, 0xFA), Inches(0.04))
    add_multiline_text(slide, Inches(9.3), Inches(1.75), Inches(3.4), Inches(2.5), [
        {"text": "Round 2+ — Deep Cascade", "size": 16, "color": RGBColor(0xA7, 0x8B, 0xFA), "bold": True, "space_after": 8},
        {"text": "📦 FedEx — logistics network strained", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "🏦 JPMorgan — loan exposure rises", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "🌾 Tyson Foods — supply chain broken", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "🏥 Tenet Healthcare — overwhelmed", "size": 12, "color": LIGHT_GRAY, "space_after": 4},
        {"text": "Losses amplify 2-5x beyond direct", "size": 11, "color": ACCENT_AMBER, "bold": True, "space_after": 4},
        {"text": "Damage decays 40% per round", "size": 11, "color": DIM_GRAY, "space_after": 0},
    ])

    # Bottom: Results summary
    results_shape = add_shape(slide, Inches(0.5), Inches(4.7), Inches(12.3), Inches(2.5),
                              RGBColor(0x0F, 0x23, 0x3D))
    add_text_box(slide, Inches(0.8), Inches(4.85), Inches(3), Inches(0.3),
                 "SIMULATION RESULTS", font_size=11, color=ACCENT_CYAN, bold=True)

    # Result metrics
    metrics = [
        ("$67.5M", "Direct Loss", ACCENT_RED),
        ("$142.3M", "Cascade Loss", ACCENT_AMBER),
        ("$209.8M", "Total Loss", WHITE),
        ("211%", "Cascade Amplification", RGBColor(0xA7, 0x8B, 0xFA)),
        ("15", "Companies Affected", ACCENT_CYAN),
    ]
    mx = Inches(0.8)
    for val, label, color in metrics:
        add_text_box(slide, mx, Inches(5.2), Inches(2), Inches(0.35),
                     val, font_size=24, color=color, bold=True, alignment=PP_ALIGN.CENTER)
        add_text_box(slide, mx, Inches(5.6), Inches(2), Inches(0.25),
                     label, font_size=10, color=MID_GRAY, alignment=PP_ALIGN.CENTER)
        mx += Inches(2.4)

    # Narrative preview
    add_text_box(slide, Inches(0.8), Inches(6.1), Inches(11.5), Inches(0.9),
                 "📝 AI Narrative: \"A high-severity hurricane event impacting Texas has triggered portfolio losses "
                 "totaling $209.8M. Direct impact accounts for $67.5M, while cascade effects through supply chain "
                 "disruptions, insurance dependencies, and grid failures added $142.3M (68% of total losses). "
                 "The most affected entities are ExxonMobil, CenterPoint Energy, and Valero Energy...\"",
                 font_size=11, color=DIM_GRAY)


def build_visualization_slide(prs):
    """Slide 8: Frontend Visualization — what users see."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_CYAN, Inches(0.04))

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(3), Inches(0.3),
                 "USER EXPERIENCE", font_size=11, color=ACCENT_CYAN, bold=True)

    add_text_box(slide, Inches(0.8), Inches(0.7), Inches(10), Inches(0.6),
                 "Premium Interactive Dashboard",
                 font_size=30, color=WHITE, bold=True)

    add_text_box(slide, Inches(0.8), Inches(1.3), Inches(10), Inches(0.4),
                 "Dark-themed, responsive interface with real-time simulation feedback and rich visualizations.",
                 font_size=14, color=LIGHT_GRAY)

    # Agent Network Graph card
    net_card = add_shape(slide, Inches(0.5), Inches(1.9), Inches(6.2), Inches(3.5), CARD_BG)
    add_accent_line(slide, Inches(0.5), Inches(1.9), Inches(6.2), ACCENT_CYAN, Inches(0.03))
    add_text_box(slide, Inches(0.7), Inches(2.05), Inches(5), Inches(0.3),
                 "🕸️  Agent Network Graph", font_size=16, color=WHITE, bold=True)

    net_features = [
        "• Canvas-based force-directed graph (not SVG — fast)",
        "• Nodes = companies, sized by loss percentage",
        "• Color-coded by sector (Energy, Utilities, Tech...)",
        "• Dashed rings = your portfolio holdings",
        "• Click any node → see cascade chain details",
        "• Hover highlights connected dependencies",
        "• Animated flow particles on focused edges",
        "• Physics simulation clusters nodes by sector",
    ]
    ny = Inches(2.5)
    for feat in net_features:
        add_text_box(slide, Inches(0.7), ny, Inches(5.8), Inches(0.28),
                     feat, font_size=11, color=LIGHT_GRAY)
        ny += Inches(0.3)

    # Right side: Other panels
    panels = [
        ("⏱️  Cascade Timeline", ACCENT_TEAL, [
            "Events grouped by round",
            "Source → Target with loss amounts",
            "Severity color coding",
        ]),
        ("📊  Affected Entities Table", ACCENT_AMBER, [
            "Sortable by loss, sector, region",
            "Portfolio vs dependency tagging",
            "Risk score filtering",
        ]),
        ("🛡️  Recommendations Panel", RGBColor(0xA7, 0x8B, 0xFA), [
            "Categorized: rebalance, hedge, exit",
            "Priority levels (high/medium/low)",
            "Affected entities per action",
        ]),
    ]

    py = Inches(1.9)
    for title, color, features in panels:
        panel = add_shape(slide, Inches(6.9), py, Inches(5.9), Inches(1.1), CARD_BG)
        add_accent_line(slide, Inches(6.9), py, Inches(0.04), color, Inches(1.1))
        add_text_box(slide, Inches(7.15), py + Inches(0.08),
                     Inches(5), Inches(0.3), title,
                     font_size=13, color=WHITE, bold=True)
        fx = Inches(7.15)
        for j, f in enumerate(features):
            add_text_box(slide, fx, py + Inches(0.42) + j * Inches(0.22),
                         Inches(5.5), Inches(0.22), f"• {f}",
                         font_size=10, color=MID_GRAY)
        py += Inches(1.2)

    # Bottom: Tech highlights
    tech_bar = add_shape(slide, Inches(0.5), Inches(5.7), Inches(12.3), Inches(1.5),
                         RGBColor(0x0F, 0x23, 0x3D))
    add_text_box(slide, Inches(0.8), Inches(5.85), Inches(3), Inches(0.3),
                 "FRONTEND TECH STACK", font_size=11, color=MID_GRAY, bold=True)

    techs = [
        ("Next.js 15", "App Router, SSR", ACCENT_CYAN),
        ("TypeScript", "Full type safety", ACCENT_TEAL),
        ("Tailwind CSS", "Utility-first styling", ACCENT_GREEN),
        ("Framer Motion", "Smooth animations", ACCENT_AMBER),
        ("Zustand", "State management", RGBColor(0xA7, 0x8B, 0xFA)),
        ("Canvas API", "Network rendering", ACCENT_RED),
    ]
    tx = Inches(0.8)
    for name, desc, color in techs:
        add_text_box(slide, tx, Inches(6.2), Inches(1.8), Inches(0.25),
                     name, font_size=13, color=color, bold=True)
        add_text_box(slide, tx, Inches(6.48), Inches(1.8), Inches(0.25),
                     desc, font_size=10, color=DIM_GRAY)
        tx += Inches(2.0)


def build_smart_features_slide(prs):
    """Slide 9: Smart features — graceful degradation, dynamic agents, etc."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BG)
    add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_GREEN, Inches(0.04))

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(3), Inches(0.3),
                 "SMART FEATURES", font_size=11, color=ACCENT_GREEN, bold=True)

    add_text_box(slide, Inches(0.8), Inches(0.7), Inches(10), Inches(0.6),
                 "Built to Be Resilient, Just Like Your Portfolio Should Be",
                 font_size=30, color=WHITE, bold=True)

    features = [
        {
            "icon": "🎯", "title": "Dynamic Agent Generation",
            "desc": "The platform doesn't just analyze your portfolio holdings. It dynamically generates "
                    "real companies relevant to the disaster — a Texas hurricane creates ExxonMobil, ERCOT, "
                    "and CenterPoint agents even if they're not in your portfolio, because they're part of "
                    "the dependency chain that affects your holdings.",
            "color": ACCENT_CYAN,
        },
        {
            "icon": "🔄", "title": "Graceful Degradation",
            "desc": "No API key? No problem. The platform works with rule-based fallback for event parsing, "
                    "cascade math, and narrative generation. With a free Gemini key, you get LLM-powered "
                    "understanding. Rate-limited? It automatically switches to rules mid-simulation.",
            "color": ACCENT_TEAL,
        },
        {
            "icon": "🌐", "title": "200+ Real Company Registry",
            "desc": "A curated database of real companies organized by sector and US region. Energy companies "
                    "in Texas, utilities in California, agriculture in Iowa — the simulation uses actual "
                    "companies with realistic market values and dependency profiles.",
            "color": ACCENT_GREEN,
        },
        {
            "icon": "💬", "title": "Interactive Agent Cascade",
            "desc": "Agents don't reason in isolation. Agent A's reaction to grid failure feeds into Agent B's "
                    "reasoning about supply chain disruption — in the same round. This creates emergent cascade "
                    "behavior that mirrors real-world interconnected crises.",
            "color": ACCENT_AMBER,
        },
        {
            "icon": "📈", "title": "Cascade Amplification Tracking",
            "desc": "The platform quantifies how much worse cascade effects make things. A 211% amplification "
                    "means indirect losses are more than double the direct damage — critical information for "
                    "risk managers who need to understand systemic exposure.",
            "color": RGBColor(0xA7, 0x8B, 0xFA),
        },
    ]

    y = Inches(1.5)
    for f in features:
        row = add_shape(slide, Inches(0.5), y, Inches(12.3), Inches(1.05), CARD_BG)
        add_accent_line(slide, Inches(0.5), y, Inches(0.04), f["color"], Inches(1.05))
        add_text_box(slide, Inches(0.75), y + Inches(0.1),
                     Inches(0.5), Inches(0.35), f["icon"],
                     font_size=18, color=f["color"])
        add_text_box(slide, Inches(1.3), y + Inches(0.1),
                     Inches(3), Inches(0.3), f["title"],
                     font_size=15, color=WHITE, bold=True)
        add_text_box(slide, Inches(1.3), y + Inches(0.45),
                     Inches(11.2), Inches(0.55), f["desc"],
                     font_size=11, color=LIGHT_GRAY)
        y += Inches(1.15)
